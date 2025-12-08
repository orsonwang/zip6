# -*- coding: utf-8 -*-
import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')
os.chdir(r'd:\data\Working\zip6')

files = [
    '中英文街路名稱對照檔1130401.xlsx',
    '鄉鎮市區中英對照.xlsx',
    '村里文字巷中英對照.xlsx'
]

for f in files:
    print(f'=== {f} ===')
    wb = openpyxl.load_workbook(f, read_only=True)
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f'Sheet: {sheet_name}')
        for i, row in enumerate(sheet.iter_rows(max_row=6, values_only=True)):
            print(f'  Row {i+1}: {row}')
    wb.close()
    print()
